# -*- coding: utf-8 -*-
"""Embed the 4 required videos into TechGAR_NCKH_13slides.pptx via python-pptx add_movie.
Positions match the SVG poster-image frames (viewBox 1280x720 -> EMU x9525)."""
from pptx import Presentation
from pptx.util import Emu

PPTX = r"D:\TechGar2\docs\TechGAR_NCKH_13slides.pptx"
VID = r"D:\TechGar2\slides_nckh\build\techgar_nckh\video"
IMG = r"D:\TechGar2\slides_nckh\build\techgar_nckh\images"
PX = 9525  # EMU per px at 1280x720 -> 13.333in x 7.5in

JOBS = [
    # (slide_index, video, poster, x, y, w, h) in px
    (0,  "clip_s01_overview.mp4",      "poster_s01.jpg", 560, 478, 360, 203),
    (6,  "demo_overlap_split.mp4",     "poster_s07.jpg",  64, 202, 760, 214),
    (10, "clip_s11_park_cycle.mp4",    "poster_s11.jpg", 952, 246, 264, 149),
    (11, "clip_s12_failure_depart.mp4","poster_s12.jpg", 672, 328, 238, 134),
]

prs = Presentation(PPTX)
for idx, vid, poster, x, y, w, h in JOBS:
    slide = prs.slides[idx]
    shp = slide.shapes.add_movie(
        f"{VID}\\{vid}",
        Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX),
        poster_frame_image=f"{IMG}\\{poster}",
        mime_type="video/mp4",
    )
    print(f"slide {idx+1}: {vid} -> shape id={shp.shape_id} name={shp.name}")
prs.save(PPTX)
print("saved:", PPTX)
